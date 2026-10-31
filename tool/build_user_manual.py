"""Gerçek mobil ekran görüntülerinden sade Word kullanım kılavuzu üretir."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SCREENS = ROOT / 'docs' / 'kilavuz' / 'ekranlar'
OUTPUT = ROOT / 'docs' / 'UBSoft_Yurt_Kontrol_Kullanim_Kilavuzu.docx'

# Başlık, kısa açıklama, gerçek uygulama görüntüsü ve kullanıcı adımları.
PAGES = [
 ('Uygulamaya giriş', 'Yurt yönetiminin verdiği e-posta ve şifreyle giriş yapın. Hesabınızın yetkisine göre yönetici veya öğrenci ekranı açılır.', ['01-giris.jpg'], ['E-posta ve şifre alanlarını doldurun, Giriş Yap düğmesine dokunun.', 'Şifrenizi unuttuysanız Şifremi unuttum seçeneğini kullanın.', 'Giriş bilgilerinizi başka kişilerle paylaşmayın.']),
 ('Mobil menü', 'Telefonda temel bölümler alt çubuktadır. Diğer yönetim ekranlarına Menü üzerinden ulaşılır.', ['29-menu.jpg'], ['Yurt Durumu, Geçmiş, Noktalar veya İzinler sekmesini açın.', 'Raporlar, Öğrenciler, Ayarlar ve Profil için Menü düğmesine dokunun.', 'Yetkiniz olmayan bölümler menüde görünmeyebilir.']),
 ('Yurt durumunu izleme', 'Bu ekran aktif öğrencilerin son kayıtlarına göre yurtta bulunma durumunu gösterir. Kartlara dokunarak listeyi daraltabilirsiniz.', ['02-yurt-durumu.jpg'], ['İçeride, son geçerli kaydı giriş olan öğrencileri gösterir.', 'Dışarıda, son geçerli kaydı çıkış olan öğrencileri gösterir; izinliler de bu sayıya dahildir.', 'Dönüşü geciken, dışarıdakiler içinde dönüş süresi geçmiş olan gruptur. Sayıları toplarken ikinci kez eklemeyin.']),
 ('Başlangıç durumunu belirleme', 'Başlangıç kaydı yok ifadesi, öğrencinin içeride veya dışarıda olduğunun henüz kayıtla belirlenmediğini anlatır.', ['28-baslangic-durumu.jpg'], ['Yurt Durumu ekranındaki İncele seçeneğine dokunun.', 'Öğrencinin gerçek durumunu doğrulayın, içeride veya dışarıda seçin ve gerekçe yazın.', 'Kaydet ile başlangıç durumu oluşturulur. Bu işlem normal QR okutmasının yerine sürekli kullanılmamalıdır.']),
 ('Öğrencileri filtreleme', 'Öğrenciyi adı, numarası veya odasıyla arayabilirsiniz. Bulunma durumu ve inceleme seçimi birlikte kullanılabilir.', ['03-gecikenler.jpg'], ['Dönüşü geciken kartına dokununca ilgili öğrenciler listelenir.', 'Öğrenci kartında son hareket zamanı ve dışarıda geçen süre görünür.', 'Yeni bir incelemeye başlamak için filtre panelindeki Sıfırla seçeneğini kullanın.']),
 ('Hareket geçmişi', 'Seçilen tarihteki kayıtlar öğrenci adları altında toplanır. Her öğrenci için hareket sayısı görünür.', ['04-gecmis.jpg'], ['Üstteki tarih alanından incelemek istediğiniz günü seçin.', 'Öğrencinin adına dokunarak gün satırını açın.', 'Bir öğrencinin kaydı görünmüyorsa seçilen tarihte hareketi bulunmayabilir.']),
 ('Aynı gün birden çok giriş ve çıkış', 'Bir güne dokunduğunuzda o günün bütün hareketleri ayrı satırlarda açılır. Yeşil giriş, kırmızı çıkış anlamına gelir.', ['05-coklu-hareket.jpg'], ['Saatleri sırayla okuyun. Bir öğrenci aynı gün birden çok kez çıkıp dönebilir.', 'Gün başlığındaki özet saatler ilk girişi ve son çıkışı gösterir; bütün hareketler alt satırlardadır.', 'Yanlış kayıtta yetkili yönetici işlem menüsünden gerekçeli geçersiz kılma yapabilir. Doğru kayıtları silerek durum düzeltmeyin.']),
 ('Rapor hazırlama', 'Raporlar bölümünde tarih aralığı, öğrenci ve hareket türü seçerek kayıtları inceleyebilirsiniz.', ['11-raporlar.jpg'], ['Tarih aralığını ve gerekiyorsa öğrenciyi seçin.', 'Hareket filtresiyle giriş, çıkış veya uyarılı kayıtları daraltın.', 'PDF indir düğmesi seçili görünümün raporunu hazırlar. İndirdiğiniz dosyayı yalnız yetkili kişilerle paylaşın.']),
 ('Raporda öğrenci seçme', 'Öğrenci alanı arama yapılabilen bir liste açar. Seçtiğiniz öğrencinin kayıtları günlere ayrılır.', ['12-ogrenci-secimi.jpg'], ['Arama alanına öğrencinin adını yazın ve adına dokunun.', 'Gün satırını açarak yeşil ve kırmızı hareketleri saatleriyle okuyun.', 'Bütün öğrencileri yeniden görmek için Tüm öğrenciler seçeneğini kullanın.']),
 ('QR noktalarını yönetme', 'Her QR noktası giriş veya çıkış için tanımlanır. Girişte ve çıkışta farklı QR kodları kullanılmalıdır.', ['06-noktalar.jpg'], ['Yeni ile yeni nokta oluşturun; listedeki noktaya dokunarak düzenleyin.', 'Noktanın yanındaki QR simgesi etiketi açar.', 'Konum ve yarıçap değerlerini gerçek kapı konumuna göre belirleyin.']),
 ('Giriş ve çıkış QR etiketleri', 'Giriş etiketi giriş, çıkış etiketi çıkış kaydı oluşturur. Etiketleri doğru kapıya görünür şekilde yerleştirin.', ['07-giris-qr.jpg','32-cikis-qr.jpg'], ['Etiketin üzerinde nokta adını ve yönünü kontrol edin.', 'Tarayıcının yazdırma özelliğiyle etiketi hazırlayın.', 'Kod değiştirilirse eski etiket çalışmaz; yeni etiketi yazdırın.']),
 ('QR noktası bilgilerini düzenleme', 'Nokta adı, kodu, yönü, koordinatı ve yarıçapı bu formdan düzenlenir. Yeni nokta oluştururken de aynı bilgiler kullanılır.', ['08-nokta-duzenle.jpg'], ['Giriş veya Çıkış yönünü doğru seçin; aynı kodu iki yön için kullanmayın.', 'Formu aşağı kaydırarak koordinat ve yarıçap alanlarını doldurun.', 'Değişiklikleri kaydet düğmesine dokunun. Etkin olmayan noktadan okutma yapılamaz.']),
 ('İzinleri izleme', 'İzinler ekranında öğrencilerin izin aralıkları ve beklenen dönüş saatleri görünür. Aktif, planlanmış, sona ermiş ve iptal edilmiş izinler ayırt edilir.', ['09-izinler.jpg'], ['Öğrenci adı, başlangıç zamanı ve dönüş zamanını birlikte kontrol edin.', 'İzin, öğrenciyi otomatik olarak dışarıda veya içeride yapmaz; hareket için QR okutulur.', 'Geçerli bir izin, kapsadığı zamanda dönüş gecikmesi değerlendirmesini etkiler.']),
 ('Öğrenci izni tanımlama', 'İzin ekleme işlemini yetkili yönetici yapar. Başlangıç ve dönüş saatleri açıkça belirtilmelidir.', ['10-izin-formu.jpg'], ['Öğrenciyi seçin, başlangıç ve dönüş zamanını girin.', 'Kısa ve gerekli bir açıklama yazın, Kaydet düğmesine dokunun.', 'Mevcut aktif veya gelecek izin varken yenisi reddedilebilir. Değişiklik gerekiyorsa mevcut izni kontrol edin.']),
 ('Öğrenci listesini yönetme', 'Öğrenciler bölümünde yurda bağlı kullanıcılar ve rolleri görünür. Yönetici hesapları da bu listede yer alabilir.', ['14-ogrenciler.jpg'], ['Yeni ile öğrenci veya yetkili yönetici hesabı oluşturun.', 'Kayıtlı kişiye dokunarak bilgilerini açın.', 'Ekrandaki kullanıcı sayısı yönetici hesaplarını da içerdiği için aktif öğrenci sayısından farklı olabilir.']),
 ('Yeni öğrenci ekleme', 'Öğrencinin girişte kullanacağı e-posta ve geçici şifreyi tanımlayın. Öğrenci numarası ve oda bilgisi aramayı kolaylaştırır.', ['15-yeni-ogrenci.jpg'], ['Ad soyad, e-posta ve geçici şifreyi girin; oda ve öğrenci numarasını ekleyin.', 'Formu aşağı kaydırın ve Öğrenci rolünü kontrol edin.', 'Öğrenciyi ekle düğmesine dokunun. Giriş bilgilerini öğrenciye güvenli bir yolla iletin.']),
 ('Hesap bilgileri ve hesap durumu', 'Kayıtlı kişinin oda, açıklama ve rol bilgileri düzenlenebilir. Hesabı kapatmak, geçmiş hareketleri ortadan kaldırmaz.', ['31-ogrenci-yonetim.jpg'], ['Bilgileri değiştirdikten sonra Değişiklikleri kaydet düğmesine dokunun.', 'Hesap etkin seçeneği kapatılırsa kişi uygulamaya giriş yapamaz.', 'Şifre sıfırlama bağlantısı gönderilebilir. Öğrenciyi sil işlemi listeden kaldırır; geçmiş hareket kayıtları korunur.']),
 ('Yurt saatlerini ayarlama', 'Genel saatler, haftanın günleri ve özel tarihler için farklı kurallar tanımlanabilir. Saatler Türkiye saatine göre değerlendirilir.', ['16-ayarlar.jpg'], ['Genel saatler alanını açın; en erken çıkış ve son dönüş saatlerini belirleyin.', 'Gerekirse haftanın bir günü veya özel tarih için farklı kural seçin.', 'Formun tamamını kontrol edip Ayarları kaydet düğmesine dokunun.']),
 ('Saat kuralını düzenleme', 'Saat kuralları erken veya geç hareketler için uyarı üretir. Hareketin kaydedilmesini tek başına engellemez.', ['17-saat-kurali.jpg'], ['En erken çıkış ve en geç dönüş saatini seçin.', 'Gerekliyse en erken giriş kuralını da açıp saatini belirleyin.', 'Uygula ile pencereyi kapatın, ardından ana formdaki Ayarları kaydet düğmesini kullanın.']),
 ('QR ve konum ayarları', 'Ayarlar formunun alt bölümünde konum zorunluluğu ve nokta dışındaki okutmaların nasıl ele alınacağı belirlenir.', ['30-konum-ayarlari.jpg'], ['Konum zorunluysa konum alınamayan okutma tamamlanamayabilir.', 'Nokta dışındaki okutmaları reddet kapalıysa kayıt konum uyarısıyla alınabilir.', 'Konum yalnız okutma sırasında ölçülür. İçeride veya dışarıda bilgisi son harekete dayanır; sürekli konum takibi değildir.']),
 ('Profil ve oturum', 'Profil ekranında hesabınızın rolünü, yurdunuzu ve son giriş zamanını görebilirsiniz. Öğrenci profilinde oda ve öğrenci numarası da yer alır.', ['18-profil.jpg'], ['Şifremi değiştir ile şifre formunu açın.', 'Aydınlatma metninden uygulamanın veri kullanımını okuyun.', 'Paylaşılan cihaz kullanıyorsanız işiniz bitince Çıkış Yap düğmesine dokunun ve çıkışı onaylayın.']),
 ('Şifre değiştirme', 'Şifre değişikliği için mevcut şifre ve yeni şifrenin iki kez girilmesi gerekir.', ['19-sifre.jpg'], ['Mevcut şifrenizi yazın.', 'Yeni şifreyi iki alana aynı şekilde girin ve Şifreyi değiştir düğmesine dokunun.', 'Şifrenizi hatırlamıyorsanız giriş ekranındaki Şifremi unuttum seçeneğini kullanın.']),
 ('Şifremi unuttum', 'Bu ekran kayıtlı e-posta adresine şifre sıfırlama bağlantısı istemek için kullanılır.', ['27-sifremi-unuttum.jpg'], ['Hesabınızda kayıtlı e-posta adresini girin.', 'Bağlantı Gönder düğmesine dokunun ve e-posta kutunuzu kontrol edin.', 'E-posta ulaşmazsa gereksiz posta klasörüne bakın; adresinizi yurt yönetimiyle doğrulayın.']),
 ('Aydınlatma metni', 'Uygulamada tutulan kimlik, hareket, konum ve izin bilgilerinin açıklamaları bu ekranda yer alır.', ['20-kvkk.jpg'], ['Metni aşağı kaydırarak bütün bölümleri okuyun.', 'Saklama süresi veya veri talepleri için yurt yönetimine başvurun.', 'Ekran görüntülerini ve raporları herkesin görebileceği alanlarda paylaşmayın.']),
 ('Öğrencinin Yurdum ekranı', 'Öğrenci giriş yaptığında kendi son durumunu ve bugünkü hareketlerini görür. Başka öğrencilerin kayıtları gösterilmez.', ['21-ogrenci-yurdum.jpg'], ['Yurttan çıkarken çıkış QR kodunu, dönerken giriş QR kodunu okutun.', 'Düğmedeki yön ve ekrandaki son durum size yol gösterir.', 'Kayıt sonrası yeni hareketin listede göründüğünü kontrol edin. Başarı bildirimi olmadan işlemi tamamlandı saymayın.']),
 ('Öğrencinin kendi geçmişi', 'Geçmiş bölümünde öğrenci kendi kayıtlarını tarih aralığı seçerek inceleyebilir.', ['22-ogrenci-gecmis.jpg'], ['Tarih aralığını seçin ve incelemek istediğiniz güne dokunun.', 'Yeşil satırlar girişleri, kırmızı satırlar çıkışları gösterir.', 'Eksik veya yanlış kayıt fark ederseniz yurt yönetimine bildirin.']),
 ('Öğrencinin izin bilgileri', 'Öğrenci yalnız kendisi için tanımlanmış izinleri ve dönüş zamanını görür. Bu ekrandan kendi iznini onaylayamaz.', ['25-ogrenci-izin.jpg'], ['İzin tarihlerini ve dönüş saatini kontrol edin.', 'İzinli çıkışta da çıkış QR kodunu, dönüşte giriş QR kodunu okutun.', 'İzin değişikliği gerekiyorsa yurt yönetimine başvurun.']),
 ('QR tarama ve kamera sorunu', 'QR tarama ekranında kapıdaki kodu çerçeve içine getirin. Bu görüntü, kamera açılamadığında uygulamanın gösterdiği yardım durumudur.', ['23-qr-tarama.jpg'], ['Kamera açılamadı yazıyorsa tarayıcının site izinlerinden kamera erişimini kontrol edin.', 'Konum isteniyorsa cihaz ve tarayıcı konum ayarlarını da kontrol edin.', 'Okutma tamamlanmazsa yönetime haber verin. Bu kılavuzdaki kamera görüntüsü başarılı bir fiziksel okutma örneği değildir.']),
]

def build():
    doc=Document()
    section=doc.sections[0]
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=Cm(1.45);section.bottom_margin=Cm(1.45)
    section.left_margin=Cm(1.8);section.right_margin=Cm(1.8)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
        style=doc.styles[name];style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.space_after=Pt(6)
    doc.styles['Normal'].font.size=Pt(10.5)
    doc.styles['Normal'].paragraph_format.line_spacing=1.08
    doc.styles['Title'].font.size=Pt(29)
    doc.styles['Heading 1'].font.size=Pt(19)
    doc.styles['Heading 1'].paragraph_format.space_before=Pt(0)
    # Word şablonundaki başlık alt çizgisini miras alma.
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    doc.styles['Caption'].font.size=Pt(9)
    doc.core_properties.title='UBSoft Yurt Kontrol Kullanım Kılavuzu'
    doc.core_properties.author='UBS Elektronik ve Güvenlik Sistemleri'
    doc.core_properties.subject='Yurt yöneticileri ve öğrenciler için kullanım adımları'
    doc.core_properties.comments=''
    doc.add_paragraph('UBSoft Yurt Kontrol Kullanım Kılavuzu','Title')
    doc.add_paragraph('Yurt yönetimi ve öğrenciler için mobil kullanım','Subtitle')
    doc.add_paragraph('Sürüm 2.1.0   |   31 Ekim 2026')
    doc.add_paragraph('Bu kılavuz; öğrencilerin giriş ve çıkışlarını kaydetmesini, yurt yönetiminin durumları izlemesini, izinleri düzenlemesini ve hareket raporlarını incelemesini anlatır. Her bölümde uygulamanın mobil görünümü ve kısa kullanım adımları yer alır.')
    doc.add_heading('İlk kullanım sırası',1)
    for x in ['Yönetici hesabıyla giriş yapın ve yurdun saat kurallarını kontrol edin.','Giriş ve çıkış için ayrı QR noktalarını hazırlayın.','Öğrenci hesaplarını, oda ve öğrenci numarası bilgilerini ekleyin.','Öğrenciyle birlikte bir çıkış ve bir giriş kaydı deneyin.','Durum ekranını, izinleri ve raporları kontrol edin.']:
        doc.add_paragraph(x,'List Number')
    doc.add_heading('Durumları doğru okuyun',1)
    for title,text in [('İçeride','Son geçerli hareket giriş kaydıdır.'),('Dışarıda','Son geçerli hareket çıkış kaydıdır.'),('Dönüşü geciken','Dışarıda olup geçerli dönüş süresini aşan öğrencidir.'),('İzinli','Belirlenen aralıkta geçerli izni vardır; fiziksel durumu yine QR kaydı belirler.'),('Başlangıç kaydı yok','İçeride veya dışarıda olduğu henüz kayıtla belirlenmemiştir.')]:
        p=doc.add_paragraph();p.add_run(title+'  ').bold=True;p.add_run(text)
    doc.add_paragraph('Uygulama adresi  https://ubsoft-pdks.web.app')
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=footer.add_run('UBSoft Yurt Kontrol   ');run.font.size=Pt(8)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    for index,(title,intro,images,steps) in enumerate(PAGES,1):
        doc.add_page_break();doc.add_heading(title,1)
        doc.add_paragraph(intro)
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after=Pt(3)
        for n,image in enumerate(images):
            if n:p.add_run('   ')
            picture=p.add_run().add_picture(str(SCREENS/image),width=Cm(7.65 if len(images)==1 else 7.45))
            props=picture._inline.docPr;props.set('descr',title+' ekranının uygulama içinden alınmış mobil görüntüsü')
        cap=doc.add_paragraph(f'Ekran {index}   {title}','Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for n,text in enumerate(steps,1):
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3)
            p.add_run(f'{n}. ').bold=True;p.add_run(text)
    doc.add_page_break();doc.add_heading('Günlük kontrol ve sık karşılaşılan durumlar',1)
    for title,text in [
      ('Günlük kontrol','Dönüşü gecikenleri, başlangıç kaydı olmayanları ve saat uyarılarını inceleyin. İzin dönüşlerini kontrol edin. Gerekli öğrencinin gününü açarak hareket sırasını okuyun.'),
      ('Öğrenci dışarıda görünüyor ama yurtta','Önce son hareketi ve doğru QR kodunun kullanılıp kullanılmadığını kontrol edin. Unutulan okutmayı veya yanlış kaydı yönetime bildirin. Gerçek durumu görmeden kayıt üzerinde işlem yapmayın.'),
      ('Öğrenci zaten içeride veya dışarıda uyarısı alıyor','Aynı yöndeki QR kodu yeniden okutulmuş olabilir. Son hareketi inceleyin. Girişte giriş, çıkışta çıkış etiketi kullanılmalıdır.'),
      ('İnternet veya konum sorunu var','Bağlantıyı ve tarayıcı izinlerini kontrol edin. Başarılı kayıt bildirimi ve hareket satırı görünmeden kaydı tamamlanmış saymayın. Sorun sürerse yönetime bildirin.'),
      ('Raporda kayıt görünmüyor','Tarih aralığını, seçili öğrenciyi ve hareket filtresini kontrol edin. Filtreleri sıfırlayın. Başlangıç durumu doğrulaması normal giriş veya çıkış hareketi olarak değerlendirilmez.'),
      ('Sunum ve gerçek kullanım','Ekrandaki durumlar son kayıtları gösterir. Okutulmayan bir fiziksel geçiş sistem tarafından kendiliğinden bilinemez. Kurumun günlük kontrol ve eksik okutma bildirim süreci önemlidir.')]:
        doc.add_heading(title,2);doc.add_paragraph(text)
    OUTPUT.parent.mkdir(parents=True,exist_ok=True)
    doc.save(OUTPUT)
    print(f'{OUTPUT}\nBölüm sayısı: {len(PAGES)}\nEkran sayısı: {sum(len(p[2]) for p in PAGES)}')

if __name__=='__main__':build()
