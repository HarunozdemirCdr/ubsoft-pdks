# -*- coding: utf-8 -*-
"""UBSoft PDKS kullanım kılavuzunu Word belgesi olarak üretir.

Müşteriye verilecek belge: sade dil, kısa paragraf, tablo yok, teknik
ayrıntı yok. Kılavuzun kaynağı bu betiktir; `.docx` dosyası çıktısıdır.
Metni düzeltmek gerektiğinde Word'de değil burada düzeltilir, aksi hâlde
iki sürüm birbirinden ayrılır.

Çalıştırma (proje kökünde):

    pip install python-docx
    python tool/kullanim_kilavuzu.py

Çıktı: Kullanim-Kilavuzu.docx
"""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Cm

VIOLET = RGBColor(0x6D, 0x28, 0xD9)
INK = RGBColor(0x1C, 0x19, 0x17)
MUTED = RGBColor(0x78, 0x71, 0x6C)

doc = Document()

# Sayfa ve temel yazı tipi
for section in doc.sections:
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.font.color.rgb = INK
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.15


def h1(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = VIOLET
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(10)
    return p


def h2(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = INK
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    return p


def para(text, muted=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    if muted:
        r.font.size = Pt(10)
        r.font.color.rgb = MUTED
    return p


def bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(text)
    p.paragraph_format.space_after = Pt(4)
    return p


# Word'un "List Number" biçemi belge boyunca kesintisiz sayar; ikinci
# bölümdeki liste 5'ten devam eder. Numarayı elle yazmak her görüntüleyicide
# doğru görünmesini garantiler.
_step = {'n': 0}


def numbered(text, first=False):
    if first:
        _step['n'] = 0
    _step['n'] += 1
    p = doc.add_paragraph()
    r = p.add_run('%d. ' % _step['n'])
    r.bold = True
    p.add_run(text)
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_after = Pt(4)
    return p


# ─────────────────────────── Kapak ───────────────────────────
cover = doc.add_paragraph()
cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = cover.add_run('UBSoft PDKS')
r.bold = True
r.font.size = Pt(34)
r.font.color.rgb = VIOLET
cover.paragraph_format.space_before = Pt(120)
cover.paragraph_format.space_after = Pt(4)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run('Personel Devam Kontrol Sistemi')
r.font.size = Pt(15)
r.font.color.rgb = INK
sub.paragraph_format.space_after = Pt(2)

sub2 = doc.add_paragraph()
sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub2.add_run('Kullanım Kılavuzu')
r.font.size = Pt(13)
r.font.color.rgb = MUTED

foot = doc.add_paragraph()
foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = foot.add_run(
    'UBS Elektronik ve Güvenlik Sistemleri\n'
    'www.ubselektronik.com  ·  0535 812 85 36'
)
r.font.size = Pt(10)
r.font.color.rgb = MUTED
foot.paragraph_format.space_before = Pt(180)

doc.add_page_break()

# ─────────────────────────── Giriş ───────────────────────────
h1('Bu uygulama ne işe yarar?')
para(
    'Personeliniz işe geldiğinde ve işten ayrılırken kapıdaki QR kodu '
    'telefonuyla okutur. Uygulama okutma anındaki saati ve kişinin bulunduğu '
    'konumu kaydeder. Böylece kimin kaçta geldiği, kaçta çıktığı ve ne kadar '
    'çalıştığı kâğıda yazmadan, imza toplamadan kayıt altına alınır.'
)
para(
    'Konum bilgisi sayesinde "işyerinde okuttu" ile "QR kodun fotoğrafını '
    'çekip evden okuttu" birbirinden ayrılır.'
)
para(
    'Uygulama telefonda, tablette ve bilgisayarda tarayıcıdan çalışır. '
    'Kurulum, indirme ya da güncelleme derdi yoktur.'
)

h2('Nasıl girilir?')
para(
    'Tarayıcıdan uygulamanın adresine girilir, e-posta ve şifre yazılır. '
    'Adres ve hesap bilgileri kurulumda tarafınıza iletilir.'
)
para(
    'Telefonda her seferinde adres yazmamak için sayfayı ana ekrana '
    'kaydedebilirsiniz; uygulama gibi açılır.'
)

# ─────────────────────────── Personel ───────────────────────────
doc.add_page_break()
h1('Personel için')

h2('İşe gelince ve çıkarken')
numbered('Uygulamayı açın.', first=True)
numbered('Ortadaki büyük düğmeye basın: "Giriş için QR okut".')
numbered('Telefonun kamerası açılır, kapıdaki QR kodu gösterin.')
numbered('"Giriş yapıldı" yazısını ve saati gördüyseniz kayıt tamamdır.')
para(
    'Çıkarken aynı şeyi yaparsınız; düğme bu kez kendiliğinden "Çıkış için '
    'QR okut" yazar. Hangi kodu okutacağınızı düşünmenize gerek yoktur, '
    'uygulama sizin o an içeride mi dışarıda mı olduğunuzu bilir.'
)

h2('Konum izni')
para(
    'İlk okutmada tarayıcı konum izni ister, "İzin Ver" demeniz gerekir. '
    'Konumunuz yalnız okutma anında, bir kerelik ölçülür. Uygulama sizi gün '
    'boyu takip etmez, arka planda çalışmaz.'
)

h2('Ekranda ne görürsünüz?')
bullet('O günkü giriş ve çıkış saatiniz.')
bullet('Çalışma süreniz.')
bullet('İçeride mi dışarıda mı olduğunuz.')
bullet('Geçmiş günlerinizdeki giriş ve çıkış saatleriniz.')
para(
    'Giriş saatleri yeşil, çıkış saatleri kırmızı yazılır; hangisinin ne '
    'olduğunu okumadan anlarsınız.'
)

h2('Şifrenizi değiştirmek')
para(
    'Profil bölümünden şifrenizi değiştirebilirsiniz. Şifrenizi unuttuysanız '
    'giriş ekranındaki "Şifremi unuttum" bağlantısı e-posta adresinize '
    'sıfırlama bağlantısı gönderir.'
)

# ─────────────────────────── Yönetici ───────────────────────────
doc.add_page_break()
h1('Yönetici için')

h2('Bugün')
para(
    'Açılışta o günün özetini görürsünüz: kaç kişi içeride, kaç kişi işini '
    'tamamladı, kim gelmedi, dikkat isteyen kaç kayıt var. Altında personel '
    'listesi ve herkesin o günkü giriş-çıkış saati durur.'
)

h2('Geçmiş')
para(
    'Bir gün seçersiniz, o günün tüm personeli tek listede görünür: solda '
    'isim, sağda giriş ve çıkış saati. Yeşil giriş, kırmızı çıkış, gri ise o '
    'gün okutma yapılmamış demektir.'
)
para(
    'Sol ve sağ oklarla günler arasında hızlıca gezinirsiniz. Bir kişinin '
    'satırına dokunduğunuzda o günkü tüm okutmaları, hangi kapıdan ve '
    'nereden okuttuğu bilgisiyle açılır.'
)
para(
    'Bu ekran günlük göz atmak içindir: "dün kim kaçta gelmiş" sorusunun '
    'cevabı buradadır.',
    muted=True,
)

h2('Raporlar')
para(
    'Bir personelin belirli tarihler arasındaki durumunu incelemek ve '
    'yazdırmak için bu ekranı kullanırsınız.'
)
numbered('Tarih aralığını seçin (örneğin ayın 1\'i ile 31\'i arası).', first=True)
numbered('İsterseniz tek bir personel seçin, istemezseniz herkes listelenir.')
numbered(
    'Bir kişinin üzerine dokunun: o tarihler arasındaki günleri gün gün, '
    'giriş ve çıkış saatleriyle açılır.'
)
numbered('"PDF olarak indir" ile aynı dökümü yazdırabilirsiniz.')
para('Ayrıca istediğiniz günleri süzebilirsiniz:')
bullet('Geç giriş yapılan günler')
bullet('Çıkışı okutulmamış günler')
bullet('Konumu doğrulanamayan okutmalar')
bullet('Hiç gelinmeyen günler')
para(
    'Örneğin "Geç giriş" seçip ay aralığını verirseniz, o ay içinde geç '
    'gelinen bütün günlerin listesini alır ve PDF olarak yazdırırsınız.'
)

h2('QR noktaları')
para(
    'Personelin okutacağı kodları buradan oluşturursunuz. Bir noktaya isim '
    'verir (Ana Giriş, Çıkış Kapısı gibi), haritadaki konumunu girer ve '
    'yazdırılabilir QR etiketini alırsınız. Etiketi kapının yanına asmanız '
    'yeterlidir.'
)
para('Üç çeşit nokta oluşturabilirsiniz:')
bullet('Yalnız giriş için')
bullet('Yalnız çıkış için')
bullet(
    'Hem giriş hem çıkış için — tek etiket yeter, uygulama yönü kendisi '
    'anlar'
)
para(
    'Noktanın konumunu girerken "Bulunduğum konumu al" düğmesine basmanız '
    'yeterlidir; kapının önünde dururken bir kez basmanız işi bitirir.'
)

h2('Personel')
para(
    'Yeni personel eklemek için ad, e-posta ve geçici bir şifre girmeniz '
    'yeterlidir. Şifreyi personele iletirsiniz, o kendi şifresini sonradan '
    'değiştirir. İşten ayrılan personeli silmek yerine pasife alırsınız; '
    'böylece geçmiş kayıtları kaybolmaz.'
)

h2('Ayarlar')
para('Şirketinize göre üç şeyi buradan ayarlarsınız:')
bullet('Mesai başlangıç ve bitiş saati')
bullet('Geç sayılma toleransı (örneğin 15 dakika)')
bullet('Konum kuralları')
para(
    'Mesai 08:30 ve tolerans 15 dakika ise, 08:45\'ten sonra yapılan girişler '
    '"geç" olarak işaretlenir.'
)

# ─────────────────────────── SSS ───────────────────────────
doc.add_page_break()
h1('Sık sorulanlar')

h2('Personel çıkışını okutmayı unutursa ne olur?')
para(
    'O gün "Çıkış yok" olarak işaretlenir ve o gün için çalışma süresi '
    'hesaplanmaz. Uygulama bir çıkış saati uydurmaz — kaç saat çalışıldığını '
    'yalnız orada olan insanlar bilir.'
)
para(
    'Ertesi gün hiçbir sorun çıkmaz: personel normal şekilde girişini okutur, '
    'yeni gün sıfırdan başlar. Unutulan gün karışıklık yaratmaz, sadece o '
    'günün süresi eksik kalır.'
)
para(
    'Bu günleri Raporlar ekranındaki "Çıkış yok" süzgeciyle topluca '
    'görebilirsiniz.'
)

h2('Personel QR kodun fotoğrafını çekip evden okutabilir mi?')
para(
    'Okutabilir ama saklayamaz. Uygulama okutma anındaki konumu ölçer ve QR '
    'noktasının konumuyla karşılaştırır. Uzaktan yapılan okutma kayda geçer '
    'ancak işaretlenir; Raporlar ekranında "Konum uyarısı" süzgeciyle bu '
    'okutmaları görürsünüz.'
)
para(
    'İsterseniz Ayarlar\'dan katı modu açarak nokta dışından yapılan '
    'okutmaların hiç kabul edilmemesini sağlayabilirsiniz.'
)

h2('Telefonun konumu kapalıysa ne olur?')
para(
    'Kayıt yine tutulur ama "konum alınamadı" olarak işaretlenir. Bu da '
    'raporlarda görünür.'
)

h2('Yanlış bir okutma yapıldı, düzeltebilir miyim?')
para(
    'Kaydın saati değiştirilemez; değiştirilebilseydi kayıt kanıt olmaktan '
    'çıkardı. Yanlış okutmayı silebilirsiniz, o günün çalışma süresi kalan '
    'okutmalardan yeniden hesaplanır. Personel isterse yeniden okutur.'
)

h2('Personel diğer arkadaşlarının kayıtlarını görebilir mi?')
para(
    'Hayır. Personel yalnız kendi giriş ve çıkışlarını görür. Konum '
    'değerlendirmesi ve geç giriş işareti gibi bilgiler de yalnız yöneticide '
    'görünür.'
)

h2('İnternet yoksa okutma yapılabilir mi?')
para(
    'Hayır, okutma anında internet bağlantısı gerekir. Kayıt anında sunucuya '
    'yazılır ve saat sunucudan alınır; böylece telefon saati değiştirilerek '
    'erken giriş yazdırılamaz.'
)

# ─────────────────────────── Destek ───────────────────────────
doc.add_page_break()
h1('Destek')
para(
    'Aklınıza takılan her şey için bize ulaşabilirsiniz. Uygulamayla ilgili '
    'talebiniz, hata bildiriminiz ya da yeni özellik isteğiniz olursa '
    'çekinmeyin.'
)
para('UBS Elektronik ve Güvenlik Sistemleri')
para('Telefon: 0535 812 85 36')
para('E-posta: info@ubselektronik.com')
para('Web: www.ubselektronik.com')

doc.save('Kullanim-Kilavuzu.docx')
print('yazildi: Kullanim-Kilavuzu.docx')
